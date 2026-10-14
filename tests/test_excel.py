"""Unit and structural validation tests for the Excel analysis model."""

from pathlib import Path
import openpyxl
import pytest

from src.config import EXCEL_DIR

EXCEL_FILE = EXCEL_DIR / "retail_sales_analysis.xlsx"


@pytest.fixture(scope="module")
def workbook():
    """Fixture to load the generated Excel workbook."""
    if not EXCEL_FILE.exists():
        pytest.skip(f"Excel workbook not found at {EXCEL_FILE}")
    # Load with data_only=False to inspect the raw formulas
    wb = openpyxl.load_workbook(str(EXCEL_FILE), data_only=False)
    yield wb
    wb.close()


def test_excel_file_exists_and_size():
    """Verify that the Excel workbook exists and has appropriate size (> 1 MB)."""
    assert EXCEL_FILE.exists()
    assert EXCEL_FILE.stat().st_size > 1_000_000


def test_workbook_sheets(workbook):
    """Verify all 6 required sheets are present in the workbook."""
    expected_sheets = [
        "Executive Dashboard",
        "Order Data",
        "Category Analysis",
        "Discount Analysis",
        "Regional Analysis",
        "Methodology",
    ]
    for s in expected_sheets:
        assert s in workbook.sheetnames, f"Missing sheet: {s}"


def test_order_data_row_and_column_count(workbook):
    """Verify Order Data sheet contains exactly 9,995 rows (1 header + 9,994 records) and 30 columns."""
    ws = workbook["Order Data"]
    assert ws.max_row == 9995
    assert ws.max_column == 30


def test_executive_dashboard_formulas(workbook):
    """Verify Executive Dashboard uses dynamic formulas pointing to Order Data."""
    ws = workbook["Executive Dashboard"]

    # Total Sales formula
    sales_formula = str(ws["C6"].value)
    assert "SUM('Order Data'!R2:R9995)" in sales_formula or "SUM('Order Data'!" in sales_formula

    # Total Profit formula
    profit_formula = str(ws["C7"].value)
    assert "SUM('Order Data'!U2:U9995)" in profit_formula or "SUM('Order Data'!" in profit_formula

    # Margin safe division formula
    margin_formula = str(ws["C8"].value)
    assert "IF(" in margin_formula and "C7/C6" in margin_formula


def test_category_analysis_formulas(workbook):
    """Verify Category Analysis uses SUMIFS formulas across subcategories."""
    ws = workbook["Category Analysis"]
    # Check that Bookcases (row 5) has SUMIFS formula
    sales_cell = str(ws["C5"].value)
    profit_cell = str(ws["D5"].value)
    margin_cell = str(ws["E5"].value)

    assert "SUMIFS('Order Data'!R" in sales_cell
    assert "SUMIFS('Order Data'!U" in profit_cell
    assert "IF(C5=0, 0, D5/C5)" in margin_cell


def test_discount_analysis_formulas(workbook):
    """Verify Discount Analysis uses COUNTIF and SUMIF for discount bands."""
    ws = workbook["Discount Analysis"]
    items_cell = str(ws["B5"].value)
    sales_cell = str(ws["C5"].value)
    profit_cell = str(ws["D5"].value)

    assert "COUNTIF('Order Data'!Z" in items_cell
    assert "SUMIF('Order Data'!Z" in sales_cell
    assert "SUMIF('Order Data'!Z" in profit_cell
